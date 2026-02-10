#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
企业信息批量检查脚本
检查企业信息表中的各项合规状态
"""

import os
import sys
import glob
import pandas as pd
from datetime import datetime
from typing import Optional, Dict, List, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ============== 配置区域 ==============
CONFIG = {
    "base_path": os.environ.get('BASE_PATH', './Downloads'),  # Excel文件所在目录
    "today_date": os.environ.get('TODAY_DATE', ''),          # 今日日期，空则自动获取
    "previous_date": os.environ.get('PREVIOUS_DATE', ''),    # 前日日期，用于比较
    "output_report": os.environ.get('OUTPUT_REPORT', ''),    # 输出报告路径，空则自动生成
}

# 模型信息
MODEL_INFO = {
    "name": "EnterpriseChecker",
    "version": "2.0.0",
    "description": "企业信息批量检查模型"
}
# ========================================


class EnterpriseCheckerError(Exception):
    """企业检查器自定义异常"""
    pass


class EnterpriseChecker:
    """企业信息检查器"""
    
    # 必需的Excel列名
    REQUIRED_COLUMNS = [
        '查询条件',
        '登记状态',
        '营业期限',
        '注销日期',
        '近一年企业作为被告的涉诉次数'
    ]
    
    # 黑名单检查项
    BLACKLIST_ITEMS = [
        '经营异常名单',
        '是否安全生产黑名单',
        '是否涉金融黑名单',
        '是否存在税务行政处罚',
        '是否存在公示处罚信息',
        '是否存在工商行政处罚',
        '是否严重违法失信企业名单'
    ]
    
    def __init__(self, base_path: str = None):
        self.base_path = base_path or CONFIG['base_path']
        self.today = datetime.now().strftime("%Y%m%d")
        self._print_model_info()
        
    def _print_model_info(self):
        """打印模型信息"""
        print("=" * 60)
        print("当前使用的模型:")
        print(f"  名称: {MODEL_INFO['name']}")
        print(f"  版本: {MODEL_INFO['version']}")
        print(f"  描述: {MODEL_INFO['description']}")
        print("=" * 60)
        print()
        
    def find_file_by_date(self, date_str: str) -> Optional[str]:
        """根据日期查找对应的Excel文件"""
        pattern = os.path.join(self.base_path, f"企业信息批量查询列表_{date_str}*.xlsx")
        files = glob.glob(pattern)
        return files[0] if files else None
    
    def _validate_columns(self, df: pd.DataFrame, file_path: str):
        """验证Excel文件是否包含必需的列"""
        missing_columns = []
        for col in self.REQUIRED_COLUMNS:
            if col not in df.columns:
                missing_columns.append(col)
        
        if missing_columns:
            raise EnterpriseCheckerError(
                f"文件 '{file_path}' 缺少必需的列: {', '.join(missing_columns)}\n"
                f"现有列: {', '.join(df.columns.tolist())}"
            )
    
    def load_excel(self, file_path: str) -> pd.DataFrame:
        """加载Excel文件"""
        if not os.path.exists(file_path):
            raise EnterpriseCheckerError(f"文件不存在: {file_path}")
        
        try:
            df = pd.read_excel(file_path)
        except Exception as e:
            raise EnterpriseCheckerError(f"无法读取Excel文件 '{file_path}': {e}")
        
        self._validate_columns(df, file_path)
        return df
    
    def check_business_status(self, row: pd.Series) -> Tuple[bool, str]:
        """检查登记状态是否为'在营（开业）'"""
        status = str(row.get('登记状态', '')).strip()
        is_valid = status == '在营（开业）'
        return is_valid, f"登记状态: {status}" + ("" if is_valid else " [异常]")
    
    def check_blacklist_items(self, row: pd.Series) -> List[Tuple[bool, str]]:
        """检查各类黑名单状态是否都为'否'"""
        results = []
        for item in self.BLACKLIST_ITEMS:
            value = str(row.get(item, '')).strip()
            is_valid = value == '否'
            results.append((is_valid, f"{item}: {value}" + ("" if is_valid else " [异常]")))
        return results
    
    def check_lawsuit_count(self, row: pd.Series) -> Tuple[bool, str]:
        """检查近一年企业作为被告的涉诉次数是否为'0'"""
        raw_value = row.get('近一年企业作为被告的涉诉次数', '')
        
        # 处理数字类型（包括 int 和 float）
        if isinstance(raw_value, (int, float)):
            is_valid = raw_value == 0 or raw_value == 0.0
            count_str = str(int(raw_value)) if not pd.isna(raw_value) else ''
        else:
            count_str = str(raw_value).strip()
            # 支持多种格式：'0', '0次', '0 次'
            is_valid = count_str in ('0', '0次', '0 次')
        
        return is_valid, f"近一年被告涉诉次数: {count_str}" + ("" if is_valid else " [异常]")
    
    def check_cancellation_date(self, row: pd.Series) -> Tuple[bool, str]:
        """检查注销日期是否为'无'"""
        date = str(row.get('注销日期', '')).strip()
        is_valid = date == '无' or date == '' or date == 'nan' or pd.isna(row.get('注销日期'))
        return is_valid, f"注销日期: {date}" + ("" if is_valid else " [异常]")
    
    def _parse_business_period_end_date(self, period: str) -> Optional[datetime]:
        """
        解析营业期限的结束日期
        支持格式：
        - 2018-03-16 至 2038-03-15
        - 2018/03/16 至 2038/03/15
        - 2038-03-15
        - 2038/03/15
        """
        if not period or period == '无':
            return None
        
        # 提取结束日期部分
        if '至' in period:
            parts = period.split('至')
            if len(parts) >= 2:
                end_date_str = parts[-1].strip()
            else:
                return None
        else:
            end_date_str = period.strip()
        
        # 标准化日期分隔符
        end_date_str = end_date_str.replace('/', '-')
        
        # 尝试多种日期格式
        date_formats = [
            "%Y-%m-%d",
            "%Y-%m-%d %H:%M:%S",
            "%Y年%m月%d日"
        ]
        
        for fmt in date_formats:
            try:
                return datetime.strptime(end_date_str, fmt)
            except ValueError:
                continue
        
        return None
    
    def check_business_period(self, row: pd.Series) -> Tuple[bool, str]:
        """检查营业期限是否有效"""
        period_value = row.get('营业期限', '')
        period = str(period_value).strip() if not pd.isna(period_value) else ''
        
        # 精确检查"长期"：只在营业期限字段本身判断
        # 避免企业名称包含"长期"导致误判
        if period in ('长期', '长期有效', '无固定期限'):
            return True, f"营业期限: {period} [长期有效]"
        
        # 检查是否以"长期"开头或结尾（如 "至 长期"）
        if period.endswith('长期') or period.startswith('长期'):
            return True, f"营业期限: {period} [长期有效]"
        
        # 尝试解析日期
        end_date = self._parse_business_period_end_date(period)
        
        if end_date:
            today = datetime.now()
            is_valid = end_date >= today
            status = "[有效]" if is_valid else "[已过期]"
            return is_valid, f"营业期限: {period} {status}"
        
        # 无法解析时，记录警告但不判定为异常
        return True, f"营业期限: {period} [无法解析日期格式，请人工核实]"
    
    def compare_with_previous(self, current_row: pd.Series, previous_df: pd.DataFrame, 
                               enterprise_name: str) -> List[Tuple[bool, str]]:
        """与前一天数据比较，检查关键信息是否变更"""
        compare_items = ['法定代表人', '经营范围', '股东持股占比', '企业地址']
        results = []
        
        # 查找前一天的对应企业记录
        prev_row = previous_df[previous_df['查询条件'] == enterprise_name]
        
        if prev_row.empty:
            results.append((True, "前一天无此企业记录，跳过比较"))
            return results
        
        prev_row = prev_row.iloc[0]
        
        for item in compare_items:
            current_value = str(current_row.get(item, '')).strip()
            prev_value = str(prev_row.get(item, '')).strip()
            
            is_same = current_value == prev_value
            if is_same:
                results.append((True, f"{item}: 未变更"))
            else:
                results.append((False, f"{item}: 已变更\n  - 原值: {prev_value}\n  - 现值: {current_value}"))
        
        return results
    
    def check_enterprise(self, enterprise_name: str, current_df: pd.DataFrame, 
                         previous_df: Optional[pd.DataFrame] = None) -> Dict:
        """检查单个企业的所有项目"""
        result = {
            'enterprise_name': enterprise_name,
            'checks': [],
            'has_issues': False
        }
        
        # 查找当前企业记录
        row = current_df[current_df['查询条件'] == enterprise_name]
        
        if row.empty:
            result['checks'].append((False, f"未找到企业: {enterprise_name}"))
            result['has_issues'] = True
            return result
        
        row = row.iloc[0]
        
        # 1. 检查登记状态
        is_valid, msg = self.check_business_status(row)
        result['checks'].append((is_valid, msg))
        if not is_valid:
            result['has_issues'] = True
        
        # 2. 检查各类黑名单
        blacklist_results = self.check_blacklist_items(row)
        for is_valid, msg in blacklist_results:
            result['checks'].append((is_valid, msg))
            if not is_valid:
                result['has_issues'] = True
        
        # 3. 检查涉诉次数
        is_valid, msg = self.check_lawsuit_count(row)
        result['checks'].append((is_valid, msg))
        if not is_valid:
            result['has_issues'] = True
        
        # 4. 检查注销日期
        is_valid, msg = self.check_cancellation_date(row)
        result['checks'].append((is_valid, msg))
        if not is_valid:
            result['has_issues'] = True
        
        # 5. 检查营业期限
        is_valid, msg = self.check_business_period(row)
        result['checks'].append((is_valid, msg))
        if not is_valid:
            result['has_issues'] = True
        
        # 6. 与前一天比较
        if previous_df is not None:
            compare_results = self.compare_with_previous(row, previous_df, enterprise_name)
            for is_valid, msg in compare_results:
                result['checks'].append((is_valid, msg))
                if not is_valid:
                    result['has_issues'] = True
        
        return result
    
    def run_check(self, today_date: str = None, previous_date: str = None) -> List[Dict]:
        """运行检查"""
        if today_date is None:
            today_date = self.today
        
        # 查找今天的文件
        today_file = self.find_file_by_date(today_date)
        if not today_file:
            raise EnterpriseCheckerError(f"未找到日期为 {today_date} 的文件，请检查文件目录: {self.base_path}")
        
        print(f"找到今日文件: {today_file}")
        current_df = self.load_excel(today_file)
        
        # 查找前一天的文件（用于比较）
        previous_df = None
        if previous_date:
            previous_file = self.find_file_by_date(previous_date)
            if previous_file:
                print(f"找到前日文件: {previous_file}")
                previous_df = self.load_excel(previous_file)
            else:
                print(f"警告: 未找到日期为 {previous_date} 的文件，跳过比较检查")
        
        # 获取所有企业名称
        enterprises = current_df['查询条件'].dropna().unique().tolist()
        print(f"共找到 {len(enterprises)} 家企业待检查")
        
        results = []
        for enterprise in enterprises:
            result = self.check_enterprise(enterprise, current_df, previous_df)
            results.append(result)
        
        return results
    
    def print_results(self, results: List[Dict]):
        """打印检查结果"""
        print("\n" + "=" * 80)
        print("企业信息检查报告")
        print("=" * 80)
        
        issues_count = 0
        for result in results:
            print(f"\n【{result['enterprise_name']}】")
            print("-" * 40)
            
            for is_valid, msg in result['checks']:
                status = "✓" if is_valid else "✗"
                print(f"  {status} {msg}")
            
            if result['has_issues']:
                issues_count += 1
                print("  >>> 存在异常项 <<<")
        
        print("\n" + "=" * 80)
        print(f"检查完成: 共 {len(results)} 家企业，{issues_count} 家存在异常")
        print("=" * 80)
        
        return issues_count
    
    def export_report_to_excel(self, results: List[Dict], output_path: str):
        """导出检查报告到Excel文件"""
        wb = Workbook()
        ws = wb.active
        ws.title = "检查报告"
        
        # 样式定义
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 写入标题行
        headers = ["企业名称", "检查项", "检查结果", "状态"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # 写入数据
        row_num = 2
        for result in results:
            enterprise_name = result['enterprise_name']
            for is_valid, msg in result['checks']:
                # 解析检查项和结果
                if ':' in msg:
                    check_item, check_result = msg.split(':', 1)
                else:
                    check_item = msg
                    check_result = ""
                
                ws.cell(row=row_num, column=1, value=enterprise_name).border = thin_border
                ws.cell(row=row_num, column=2, value=check_item.strip()).border = thin_border
                ws.cell(row=row_num, column=3, value=check_result.strip()).border = thin_border
                
                status_cell = ws.cell(row=row_num, column=4, value="通过" if is_valid else "异常")
                status_cell.border = thin_border
                status_cell.fill = pass_fill if is_valid else fail_fill
                status_cell.alignment = Alignment(horizontal='center')
                
                row_num += 1
        
        # 调整列宽
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 10
        
        # 添加汇总sheet
        ws_summary = wb.create_sheet(title="汇总")
        ws_summary.cell(row=1, column=1, value="检查汇总").font = Font(bold=True, size=14)
        ws_summary.cell(row=3, column=1, value="检查时间:")
        ws_summary.cell(row=3, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ws_summary.cell(row=4, column=1, value="企业总数:")
        ws_summary.cell(row=4, column=2, value=len(results))
        ws_summary.cell(row=5, column=1, value="异常企业数:")
        issues_count = sum(1 for r in results if r['has_issues'])
        ws_summary.cell(row=5, column=2, value=issues_count)
        ws_summary.cell(row=6, column=1, value="正常企业数:")
        ws_summary.cell(row=6, column=2, value=len(results) - issues_count)
        
        # 异常企业列表
        if issues_count > 0:
            ws_summary.cell(row=8, column=1, value="异常企业列表:").font = Font(bold=True)
            row_num = 9
            for result in results:
                if result['has_issues']:
                    ws_summary.cell(row=row_num, column=1, value=result['enterprise_name'])
                    row_num += 1
        
        wb.save(output_path)
        print(f"\n检查报告已导出到: {output_path}")


def main():
    """主函数"""
    # 从配置区域读取参数
    base_path = CONFIG['base_path']
    today_date = CONFIG['today_date'] or datetime.now().strftime("%Y%m%d")
    previous_date = CONFIG['previous_date']
    output_report = CONFIG['output_report']
    
    print(f"配置信息:")
    print(f"  - 文件目录: {base_path}")
    print(f"  - 今日日期: {today_date}")
    if previous_date:
        print(f"  - 前日日期: {previous_date}")
    if output_report:
        print(f"  - 输出报告: {output_report}")
    print()
    
    exit_code = 0
    
    try:
        checker = EnterpriseChecker(base_path=base_path)
        results = checker.run_check(today_date=today_date, previous_date=previous_date or None)
        issues_count = checker.print_results(results)
        
        # 导出报告
        if output_report:
            checker.export_report_to_excel(results, output_report)
        else:
            # 默认导出到数据目录
            default_report = f"企业检查报告_{today_date}.xlsx"
            report_path = os.path.join(base_path, default_report)
            checker.export_report_to_excel(results, report_path)
        
        # 如果有异常，返回非零退出码
        if issues_count > 0:
            exit_code = 1
            
    except EnterpriseCheckerError as e:
        print(f"\n错误: {e}", file=sys.stderr)
        exit_code = 2
    except Exception as e:
        print(f"\n未知错误: {e}", file=sys.stderr)
        exit_code = 3
    
    sys.exit(exit_code)


if __name__ == "__main__":
    main()
